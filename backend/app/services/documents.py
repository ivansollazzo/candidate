"""Text extraction from PDF, Word and plain-text files."""

import io
from pathlib import Path

import pdfplumber
from docx import Document

from app.i18n import t

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedDocument(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text = _pdf_text(data)
    elif suffix == ".docx":
        text = _docx_text(data)
    elif suffix in {".txt", ".md"}:
        text = data.decode("utf-8", errors="replace")
    else:
        raise UnsupportedDocument(t("unsupported_format", suffix=suffix or t("unknown_format"),
                                    supported=", ".join(sorted(SUPPORTED_EXTENSIONS))))
    text = _normalize(text)
    if not text:
        raise UnsupportedDocument(t("no_text"))
    return text


# Multi-column layout detection (common in CVs with a sidebar)
MIN_GUTTER_PT = 12  # minimum gap width between columns
MAX_CROSSING_WORDS = 3  # words allowed to cross the gap (e.g. a full-width heading)
MIN_COLUMN_WORDS = 15


def _pdf_text(data: bytes) -> str:
    pages: list[str] = []
    previous_gutter: float | None = None
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            text, previous_gutter = _page_text(page, previous_gutter)
            pages.append(text)
    return "\n\n".join(pages)


def _page_text(page, previous_gutter: float | None) -> tuple[str, float | None]:
    """Read one column at a time; otherwise lines from side-by-side columns get interleaved."""
    words = page.extract_words()
    gutter = _find_gutter(page.width, words)
    if gutter is not None:
        left = sum(1 for w in words if w["x1"] <= gutter)
        right = len(words) - left
        # A page continuing the previous page's columns may have a nearly empty column
        continues = previous_gutter is not None and abs(gutter - previous_gutter) < 20
        if min(left, right) >= (1 if continues else MIN_COLUMN_WORDS) and _right_column_aligned(words, gutter):
            def column(keep_left: bool):
                return page.filter(
                    lambda o: o.get("object_type") != "char" or (((o["x0"] + o["x1"]) / 2 < gutter) == keep_left)
                )

            text = "\n\n".join(column(side).extract_text() or "" for side in (True, False))
            return text, gutter
    return page.extract_text() or "", None


def _right_column_aligned(words: list[dict], gutter: float) -> bool:
    """A real column has left-aligned lines; right-aligned dates or words don't count."""
    starts: dict[int, float] = {}
    for w in words:
        if w["x0"] >= gutter:
            row = round(w["top"] / 3)
            starts[row] = min(starts.get(row, w["x0"]), w["x0"])
    if len(starts) < 3:
        return False
    values = sorted(starts.values())
    best = max(sum(1 for v in values if abs(v - anchor) <= 4) for anchor in values)
    return best / len(values) >= 0.6


def _find_gutter(width: float, words: list[dict]) -> float | None:
    """Find the widest vertical band (in the central area) that no word crosses.

    If there is none, tolerate a few crossing words (e.g. a full-width heading).
    """
    size = int(width) + 1
    coverage = [0] * size
    for w in words:
        for x in range(max(0, int(w["x0"])), min(size, int(w["x1"]) + 1)):
            coverage[x] += 1
    for max_crossing in (0, MAX_CROSSING_WORDS):
        best: tuple[int, int] | None = None
        start: int | None = None
        for x in range(int(size * 0.15), int(size * 0.85) + 1):
            if x < int(size * 0.85) and coverage[x] <= max_crossing:
                start = x if start is None else start
            elif start is not None:
                if best is None or x - start > best[1] - best[0]:
                    best = (start, x)
                start = None
        if best and best[1] - best[0] >= MIN_GUTTER_PT:
            return (best[0] + best[1]) / 2
    return None


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _normalize(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    out: list[str] = []
    for line in lines:
        if not line and out and not out[-1]:
            continue
        out.append(line)
    return "\n".join(out).strip()
