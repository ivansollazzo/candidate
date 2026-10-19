import io

import numpy as np
import pytest
from docx import Document

from app.schemas import (
    AdviceExtraction,
    ExtractedSkill,
    HistoryExtraction,
    JobCompensationExtraction,
    JobCoreExtraction,
    JobRequirementsExtraction,
    PersonalExtraction,
    SkillsExtraction,
)
from app.services import extraction, matching

JOB_TEXT = "Cerchiamo un Backend Developer Python a Milano, RAL 35-42k. Requisiti: Python, Postgres, Docker. Gradito Kubernetes."


def fake_llm(system, user, schema):
    if schema is JobCoreExtraction:
        return JobCoreExtraction(title="Backend Developer", company="ACME", location="Milano")
    if schema is JobCompensationExtraction:
        # The model skips the salary: it must come from the deterministic parser
        assert "RAL 35-42k" in user
        return JobCompensationExtraction(benefits=["Buoni pasto"])
    if schema is JobRequirementsExtraction:
        return JobRequirementsExtraction(
            requirements=["Python", "PostgreSQL", "Docker"],
            skills=[ExtractedSkill(name="Python", importance="required"),
                    ExtractedSkill(name="Postgres", importance="required"),
                    ExtractedSkill(name="Docker", importance="required"),
                    ExtractedSkill(name="Kubernetes", importance="nice_to_have")],
        )
    if schema is PersonalExtraction:
        # The model "forgets" email and phone: they must come from the deterministic fallback
        return PersonalExtraction(full_name="Ivan Rossi", headline="Software engineer")
    if schema is HistoryExtraction:
        return HistoryExtraction(
            experiences=[{"role": "Dev", "company": "A", "start": "2020-01", "end": "2021-12"},
                         {"role": "Dev", "company": "B", "start": "2021-06", "end": "2022-12"}],
            education=[{"degree": "Laurea", "institution": "UniPA", "details": "Esame finale: Docker"}],
        )
    if schema is SkillsExtraction:
        return SkillsExtraction(skills=[{"name": "Python", "level": 3}, {"name": "PostgreSQL", "level": 2},
                                        {"name": "Docker Compose", "level": 3}, {"name": "python", "level": 9}])
    if schema is AdviceExtraction:
        return AdviceExtraction(summary="Buona compatibilità", improvements=[{"skill": "Kubernetes", "suggestion": "Corso CKAD"}])
    raise AssertionError(schema)


def fake_embed(texts):
    # Fake vectors: "docker" and "docker compose" very close, the rest orthogonal
    vocab = ["python", "postgresql", "docker", "kubernetes"]
    out = []
    for t in texts:
        v = np.array([1.0 if w in t.lower() else 0.0 for w in vocab] + [0.3 if "compose" in t.lower() else 0.0])
        out.append(v / (np.linalg.norm(v) or 1))
    return np.array(out)


@pytest.fixture(autouse=True)
def mocks(monkeypatch):
    monkeypatch.setattr(extraction, "generate_structured", fake_llm)
    monkeypatch.setattr(matching, "embed", fake_embed)


def ai_fill(client, path, **kwargs) -> dict:
    """Run an LLM fill and return its result (background tasks are synchronous in tests)."""
    r = client.post(path, **kwargs)
    assert r.status_code == 202, r.text
    task = client.get(f"/api/ai/tasks/{r.json()['id']}").json()
    assert task["status"] == "done", task
    return {**task["result"], "source_text": r.json()["source_text"]}


def test_manual_application_lifecycle(client):
    r = client.post("/api/applications", json={"title": "Backend Developer", "company": "ACME"})
    assert r.status_code == 201
    data = r.json()
    assert data["status"] == "applied" and data["applied_on"] is not None
    assert data["has_description"] is False and "description" not in data
    app_id = data["id"]

    r = client.patch(f"/api/applications/{app_id}", json={"status": "interview", "notes": "Inviata via LinkedIn"})
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "interview"
    assert body["events"][0]["type"] == "status_change"

    r = client.post("/api/events", json={"application_id": app_id, "type": "interview", "occurred_on": "2026-10-20",
                                         "starts_at": "15:00", "title": "Colloquio tecnico"})
    assert r.status_code == 201
    assert len(client.get(f"/api/applications/{app_id}").json()["events"]) == 2

    assert [a["id"] for a in client.get("/api/applications").json()] == [app_id]
    assert client.delete(f"/api/applications/{app_id}").status_code == 204


def test_ai_fills_job_posting_without_saving(client):
    fields = ai_fill(client, "/api/ai/job-posting", data={"text": JOB_TEXT})
    assert fields["title"] == "Backend Developer"
    assert (fields["salary"], fields["salary_period"], fields["salary_basis"]) == ("35.000 – 42.000 €", "annual", "gross")
    assert fields["benefits"] == ["Buoni pasto"] and fields["source_text"] == JOB_TEXT
    assert client.get("/api/applications").json() == []

    # The frontend saves the fields with the text, which stays available for refilling
    payload = {k: v for k, v in fields.items() if k != "source_text"} | {"description": JOB_TEXT}
    app = client.post("/api/applications", json=payload).json()
    assert app["has_description"] and app["skills"][0]["name"] == "Python"
    again = ai_fill(client, "/api/ai/job-posting", data={"application_id": str(app["id"])})
    assert again["company"] == "ACME"


def test_ai_reads_link_and_file(client, monkeypatch):
    from app.api import ai
    from app.services.webpage import FetchError

    monkeypatch.setattr(ai, "fetch_job_posting", lambda url: JOB_TEXT)
    assert ai_fill(client, "/api/ai/job-posting", data={"url": "https://example.com/job/1"})["source_text"] == JOB_TEXT

    def blocked(url):
        raise FetchError("Il sito ha bloccato lo scaricamento automatico")
    monkeypatch.setattr(ai, "fetch_job_posting", blocked)
    r = client.post("/api/ai/job-posting", data={"url": "https://example.com/job/2"})
    assert r.status_code == 422 and "bloccato" in r.json()["detail"]

    r = client.post("/api/ai/job-posting", files={"file": ("a.png", b"xxx", "image/png")})
    assert r.status_code == 422


def test_llm_disabled_blocks_ai_features(client):
    assert client.get("/api/settings").json()["llm_enabled"] is True
    assert client.patch("/api/settings", json={"llm_enabled": False}).json()["llm_enabled"] is False
    health = client.get("/api/health").json()
    assert health["llm"] == {"enabled": False, "hardware": None}
    assert (health["version"], health["build"]) == ("1.1", "20261019.1528")

    r = client.post("/api/ai/job-posting", data={"text": JOB_TEXT})
    assert r.status_code == 409 and "disattivate" in r.json()["detail"]
    app_id = client.post("/api/applications", json={"title": "Manuale"}).json()["id"]
    assert client.post(f"/api/applications/{app_id}/advice").status_code == 409
    assert client.post("/api/ai/cv").status_code == 409


def test_cv_upload_ai_fill_and_match(client):
    doc = Document()
    doc.add_paragraph("Ivan Rossi - Software engineer. Python, PostgreSQL, Docker Compose.")
    doc.add_paragraph("ivan.rossi@example.com · +39 333 123 4567 · github.com/ivanrossi · 2018 - 2023")
    buf = io.BytesIO()
    doc.save(buf)
    r = client.post("/api/profile/cv", files={"file": ("cv.docx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 200 and r.json()["cv_filename"] == "cv.docx"
    # Uploading alone fills nothing in
    assert r.json()["full_name"] is None

    fields = ai_fill(client, "/api/ai/cv")
    assert fields["full_name"] == "Ivan Rossi"
    assert fields["email"] == "ivan.rossi@example.com"
    assert fields["phone"] == "+39 333 123 4567"
    assert fields["links"] == [{"label": "GitHub", "url": "https://github.com/ivanrossi"}]
    assert fields["certifications"] == []
    assert fields["years_of_experience"] == 3.0  # 2020-01..2022-12, overlap counted once
    # Duplicates removed (highest level kept), levels clamped to 1-4
    assert {s["name"].lower(): s["level"] for s in fields["skills"]}["python"] == 4
    client.patch("/api/profile", json={k: v for k, v in fields.items() if k != "source_text"})

    job = ai_fill(client, "/api/ai/job-posting", data={"text": JOB_TEXT})
    app_id = client.post("/api/applications", json={"skills": job["skills"], "title": job["title"]}).json()["id"]
    match = client.get(f"/api/applications/{app_id}/match").json()
    assert {m["skill"] for m in match["matched"]} == {"Python", "Postgres", "Docker"}
    assert [m["skill"] for m in match["missing"]] == ["Kubernetes"]
    assert match["score"] == 86  # 6 of 7 points

    client.post(f"/api/applications/{app_id}/advice")
    advice = client.get(f"/api/applications/{app_id}").json()["advice"]
    assert advice["improvements"][0]["skill"] == "Kubernetes"


def test_ai_cv_requires_upload(client):
    assert client.post("/api/ai/cv").status_code == 409


def test_skill_level_affects_match(client):
    client.patch("/api/profile", json={"skills": [{"name": "Python", "level": 1}, {"name": "Docker", "level": 2}]})
    skills = [{"name": n, "importance": "required"} for n in ("Python", "Docker")]
    app_id = client.post("/api/applications", json={"title": "Dev", "skills": skills}).json()["id"]
    match = client.get(f"/api/applications/{app_id}/match").json()
    partial = {m["skill"]: m for m in match["partial"]}
    assert partial["Python"]["level"] == 1 and "base" in partial["Python"]["note"]
    assert [m["skill"] for m in match["matched"]] == ["Docker"]

    client.patch(f"/api/applications/{app_id}", json={"seniority": "Senior"})
    match = client.get(f"/api/applications/{app_id}/match").json()
    assert match["matched"] == [] and {m["skill"] for m in match["partial"]} == {"Python", "Docker"}


def test_profile_sections_editable(client):
    payload = {
        "certifications": [{"name": "AWS Certified Developer", "issuer": "Amazon", "issued_on": "2025-03"}],
        "projects": [{"name": "Roboverse", "context": "accademico", "technologies": ["ROS 2"]}],
        "languages": [{"language": "Inglese", "level": "B2"}],
    }
    profile = client.patch("/api/profile", json=payload).json()
    assert profile["certifications"][0]["issuer"] == "Amazon"
    assert profile["projects"][0]["technologies"] == ["ROS 2"]
    profile = client.patch("/api/profile", json={"certifications": []}).json()
    assert profile["certifications"] == [] and profile["projects"]


def test_settings_and_language(client):
    settings = client.get("/api/settings").json()
    assert (settings["llm_enabled"], settings["llm_accelerator"], settings["password_min_length"]) == (True, "auto", 10)
    assert "secret_key" not in settings
    assert client.patch("/api/settings", json={"llm_accelerator": "vulkan"}).json()["llm_accelerator"] == "vulkan"
    assert client.patch("/api/settings", json={"llm_accelerator": "tpu"}).status_code == 422
    assert client.patch("/api/settings", json={"session_hours": 0}).status_code == 422

    me = client.get("/api/me").json()
    assert (me["language"], me["theme"], me["onboarded"]) == ("it", "system", False)
    me = client.patch("/api/me", json={"language": "en", "theme": "dark", "onboarded": True}).json()
    assert (me["language"], me["theme"], me["onboarded"]) == ("en", "dark", True)
    assert client.patch("/api/me", json={"theme": "blue"}).status_code == 422
    assert me["application_sort"] == "applied"
    assert client.patch("/api/me", json={"application_sort": "event"}).json()["application_sort"] == "event"
    assert client.patch("/api/me", json={"application_sort": "salary"}).status_code == 422

    # Messages follow the interface language
    assert client.get("/api/applications/999", headers={"Accept-Language": "en"}).json()["detail"] == "Application not found"
    assert client.get("/api/applications/999", headers={"Accept-Language": "it-IT,it;q=0.9"}).json()["detail"] == \
        "Candidatura non trovata"


def test_status_change_stored_as_codes(client):
    app_id = client.post("/api/applications", json={"title": "Dev"}).json()["id"]
    events = client.patch(f"/api/applications/{app_id}", json={"status": "offer"}).json()["events"]
    assert events[0]["title"] == "applied>offer"


def test_scheduled_events(client):
    app_id = client.post("/api/applications", json={"title": "Backend Developer", "company": "ACME"}).json()["id"]
    client.patch(f"/api/applications/{app_id}", json={"status": "interview"})
    assert client.post("/api/events", json={"application_id": 999, "occurred_on": "2099-10-20", "title": "x"}).status_code == 404

    event = client.post("/api/events", json={"application_id": app_id, "type": "interview", "occurred_on": "2099-10-20",
                                             "starts_at": "15:30", "title": "Colloquio tecnico", "details": "Con il CTO"}).json()
    assert (event["completed"], event["starts_at"]) == (False, "15:30:00")
    client.post("/api/events", json={"application_id": app_id, "type": "follow_up", "occurred_on": "2099-10-18",
                                     "title": "Scrivere al recruiter"})

    # The agenda excludes status changes and includes the application
    agenda = client.get("/api/events").json()
    assert [e["title"] for e in agenda] == ["Scrivere al recruiter", "Colloquio tecnico"]
    assert (agenda[0]["application_title"], agenda[0]["application_company"]) == ("Backend Developer", "ACME")

    updated = client.patch(f"/api/events/{event['id']}", json={"completed": True, "details": "Andato bene"}).json()
    assert (updated["completed"], updated["details"], updated["title"]) == (True, "Andato bene", "Colloquio tecnico")
    assert client.patch(f"/api/events/{event['id']}", json={"title": None}).status_code == 422

    assert client.delete(f"/api/events/{event['id']}").status_code == 204
    assert client.delete(f"/api/events/{event['id']}").status_code == 404
    assert len(client.get("/api/events").json()) == 1
    # Status changes have already happened
    status_change = client.get(f"/api/applications/{app_id}").json()["events"][0]
    assert (status_change["type"], status_change["completed"]) == ("status_change", True)


def test_project_technologies_become_skills(client):
    client.patch("/api/profile", json={"skills": [{"name": "Python", "level": 3}]})
    profile = client.patch("/api/profile", json={"projects": [
        {"name": "Assistente", "technologies": ["Large Language Models", "python", " "]},
    ]}).json()
    # Existing skills (case-insensitive) aren't repeated; new ones come without a level
    assert [(s["name"], s["level"]) for s in profile["skills"]] == [("Python", 3), ("Large Language Models", None)]
    assert profile["skills"][1]["evidence"] == "Progetto Assistente"

    # Level stays empty until the user sets it
    skills = [profile["skills"][0], {**profile["skills"][1], "level": 2}]
    assert client.patch("/api/profile", json={"skills": skills}).json()["skills"][1]["level"] == 2


def test_profile_open_adds_missing_project_skills(client):
    from sqlmodel import Session, select

    from app.db import engine
    from app.models import Profile

    client.get("/api/profile")
    with Session(engine) as session:
        profile = session.exec(select(Profile)).one()
        profile.projects = [{"name": "Vecchio", "technologies": ["ROS 2"]}]
        session.add(profile)
        session.commit()
    assert [s["name"] for s in client.get("/api/profile").json()["skills"]] == ["ROS 2"]


def test_experience_technologies_become_skills(client):
    profile = client.patch("/api/profile", json={"experiences": [
        {"role": "Robotics engineer", "company": "Acme", "technologies": ["ROS 2"]},
    ]}).json()
    assert profile["skills"] == [{"name": "ROS 2", "category": None, "level": None,
                                  "evidence": "Esperienza Robotics engineer · Acme"}]


def test_list_reports_last_event(client):
    first = client.post("/api/applications", json={"title": "A"}).json()["id"]
    second = client.post("/api/applications", json={"title": "B"}).json()["id"]
    for day in ("2026-03-01", "2026-05-10"):
        client.post("/api/events", json={"application_id": first, "type": "interview", "occurred_on": day, "title": "x"})
    # Status changes don't count as events
    client.patch(f"/api/applications/{second}", json={"status": "inactive"})
    last = {a["id"]: a["last_event_on"] for a in client.get("/api/applications").json()}
    assert last == {first: "2026-05-10", second: None}
